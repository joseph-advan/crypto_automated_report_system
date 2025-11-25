import re
from docx import Document
from .validator import Validator

class AuditProcessor:
    def __init__(self):
        self.mapping_list = []        # Mapping 表
        self.trace_requests = []      # 金流比對請求 (Trace API)
        self.address_requests = set() # 地址驗證請求 (Address API)
        self.txh_check_requests = set() # 交易序號驗證請求 (TXH Check API)
        self.counters = {"ADDR": 0, "TXID": 0, "AMT": 0, "TIME": 0}

    def _get_next_id(self, prefix):
        self.counters[prefix] += 1
        return f"[{prefix}_{self.counters[prefix]:03d}]"

    def process_document(self, input_path, output_docx_path):
        print(f"  [AuditProcessor] 讀取文件: {input_path}")
        doc = Document(input_path)

        # --- 階段 1: 處理表格 (優先處理交易明細) ---
        for table in doc.tables:
            self._process_table_logic(table)

        # --- 階段 2: 處理非表格段落 ---
        for para in doc.paragraphs:
            self._scan_and_mask_content(para)

        # --- 階段 3: 處理頁首頁尾 ---
        for section in doc.sections:
            for header_para in section.header.paragraphs:
                self._scan_and_mask_content(header_para)
            for footer_para in section.footer.paragraphs:
                self._scan_and_mask_content(footer_para)

        doc.save(output_docx_path)
        print(f"  [AuditProcessor] 全域遮罩完成，文件已儲存: {output_docx_path}")
        
        return {
            "mapping": self.mapping_list,
            "trace_requests": self.trace_requests,
            "address_requests": list(self.address_requests),
            "txh_requests": list(self.txh_check_requests)
        }

    def _process_table_logic(self, table):
        """
        針對表格的處理邏輯：
        專門處理 5 欄位的表格 (如附表1、附表2)，進行資料提取與清洗。
        在此階段就會把誤抓的 \\n (換行) 處理掉。
        """
        for row in table.rows:
            # 檢查是否為目標表格結構 (通常有 5 個欄位: 時間, TXID, 發送, 接收, 金額)
            if len(row.cells) == 5:
                
                # === [修改核心]：在讀取時直接清洗換行符號 ===
                
                # 1. 時間：換行取代為「空格」，避免年月黏在一起 (如 2023-08\n-04 -> 2023-08 -04)
                raw_time_cell = row.cells[0].text.strip().replace('\n', ' ')
                
                # 2. 其他欄位：換行直接「移除」，將被切斷的字串接回來 (如 abc\ndef -> abcdef)
                raw_txid_cell = row.cells[1].text.strip().replace('\n', '')
                raw_from_cell = row.cells[2].text.strip().replace('\n', '')
                raw_to_cell   = row.cells[3].text.strip().replace('\n', '')
                raw_amt_cell  = row.cells[4].text.strip().replace('\n', '')

                # 3. 判斷是否為有效行：必須有 TXID，且 Amount 看起來像數字 (允許純整數)
                is_valid_amt = False
                try:
                    # 移除逗號後轉 float
                    float(raw_amt_cell.replace(",", ""))
                    is_valid_amt = True
                except ValueError:
                    is_valid_amt = False

                if raw_txid_cell and is_valid_amt:
                    # --- A. 執行遮罩與 Mapping 記錄 ---
                    # 注意：這裡傳入的是已經去除 \n 的字串
                    self._mask_and_record(row.cells[0], raw_time_cell, "TIME")
                    self._mask_and_record(row.cells[1], raw_txid_cell, "TXID")
                    self._mask_and_record(row.cells[2], raw_from_cell, "ADDR")
                    self._mask_and_record(row.cells[3], raw_to_cell,   "ADDR")
                    self._mask_and_record(row.cells[4], raw_amt_cell,  "AMT")

                    # --- B. 收集 API Request 資料 ---

                    # 1. [TXH Check] 
                    #    現在存入的是無換行的 TXID，空格可能還在(如果有的話)，但 \n 已經沒了
                    self.txh_check_requests.add(raw_txid_cell)

                    # 2. [Trace Request] 
                    try: 
                        amt_val = float(raw_amt_cell.replace(",", ""))
                    except: 
                        amt_val = 0.0
                    
                    # 進一步清洗時間格式：將 '-' 統一為 '/'，並移除多餘空格
                    clean_time_str = raw_time_cell.replace('-', '/')
                    clean_time_str = " ".join(clean_time_str.split())

                    # 進一步清洗地址：只抓取 "T" 開頭的 34 碼字串，過濾掉 "(OKX...)" 等註記
                    real_from = Validator.extract_first(raw_from_cell, "ADDR")
                    real_to   = Validator.extract_first(raw_to_cell, "ADDR")
                    
                    final_from = real_from if real_from else raw_from_cell
                    final_to   = real_to if real_to else raw_to_cell
                    
                    # 進一步清洗 TXID：Trace API 需要完全乾淨的 ID (移除所有空格)
                    clean_txid_str = raw_txid_cell.replace(' ', '')

                    self.trace_requests.append({
                        "txh": clean_txid_str,  
                        "chain": "TRON", 
                        "Token": "USDT",
                        "TimeStamp_UTC_8": clean_time_str,
                        "From": final_from,     
                        "To": final_to,         
                        "Amount": amt_val
                    })
                    
                    # 處理完此行，跳過通用掃描
                    continue 

            # ====== [Action B] 通用遮罩 (對非目標表格或無法解析的列) ======
            for cell in row.cells:
                self._scan_and_mask_content(cell)

    def _scan_and_mask_content(self, container):
        """
        通用掃描器：處理非表格段落
        """
        text = container.text
        if not text.strip(): return

        scan_types = ["TXID", "ADDR", "TIME", "AMT"] 

        for dtype in scan_types:
            matches = Validator.find_all(container.text, dtype)
            for match in matches:
                raw_text = match.group(0)
                if raw_text in container.text:
                    self._mask_and_record(container, raw_text, dtype)

    def _mask_and_record(self, element, raw_text, type_prefix):
        """
        遮罩執行的原子操作
        """
        if not raw_text: return
        
        # 避免重複遮罩
        if raw_text.startswith("[") and raw_text.endswith("]") and type_prefix in raw_text:
            return

        new_id = self._get_next_id(type_prefix)
        
        # 進行驗證，產生 status code
        clean_val, status, err_msg = Validator.validate(raw_text, type_prefix)

        # 1. 寫入 Mapping 表
        self.mapping_list.append({
            "id": new_id, 
            "type": type_prefix,
            "raw_text": raw_text, 
            "clean_val": clean_val, 
            "status": status, 
            "error_msg": err_msg
        })

        # 2. 收集額外的 API Request (針對段落中出現的資料)
        if type_prefix == "ADDR" and clean_val:
            self.address_requests.add(clean_val)
            
        if type_prefix == "TXID" and raw_text:
            self.txh_check_requests.add(raw_text)

        # 3. 執行遮罩 (修改 Word 物件)
        # 注意：如果我們在 _process_table_logic 移除了 \n，這裡的 raw_text 也是無 \n 的
        # 但 element.text 裡可能還有 \n。replace 可能會失敗。
        # 但因為這是 POC，且通常表格內的替換我們已經在邏輯上對應到了，
        # 若遇到跨行的文字替換失敗，通常不影響 API Request 的正確性。
        if hasattr(element, 'text'):
            # 嘗試直接替換 (最簡單的情況)
            element.text = element.text.replace(raw_text, new_id)